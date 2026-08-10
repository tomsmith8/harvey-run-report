#!/usr/bin/env python3
"""Stage 2.5 (optional): download source documents + the final deliverable and
convert them for the in-page viewer. Also computes rubric -> document links.

Usage: python fetch_docs.py --out out/
Requires: python-docx, openpyxl. Skips gracefully on download or parse errors.
"""
import argparse
import email
import hashlib
import email.policy
import html as H
import json
import pathlib
import re
import urllib.request


def docx_to_html(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    parts, plain = [], []
    for child in d.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or '').lower()
            plain.append(t)
            e = H.escape(t)
            if 'heading 1' in style or 'title' in style:
                parts.append(f'<h2>{e}</h2>')
            elif 'heading' in style:
                parts.append(f'<h3>{e}</h3>')
            else:
                bold = all(r.bold for r in p.runs if r.text.strip()) and any(r.text.strip() for r in p.runs)
                parts.append(f'<h4>{e}</h4>' if bold and len(t) < 90 else f'<p>{e}</p>')
        elif tag == 'tbl':
            tb = Table(child, d)
            rows = []
            for i, row in enumerate(tb.rows):
                try:
                    cells = [c.text.strip() for c in row.cells]
                except Exception:
                    continue
                plain.append(' | '.join(cells))
                tc = 'th' if i == 0 else 'td'
                rows.append('<tr>' + ''.join(f'<{tc}>{H.escape(c)}</{tc}>' for c in cells) + '</tr>')
            parts.append('<div class="tblwrap"><table>' + ''.join(rows) + '</table></div>')
    return '\n'.join(parts), '\n'.join(plain)


def eml_to_html(path):
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    body_txt = None
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == 'text/plain':
                body_txt = part.get_content()
                break
    else:
        body_txt = msg.get_content()
    body_txt = body_txt or ''
    hdr = f"From: {msg['From']}\nTo: {msg['To']}\nDate: {msg['Date']}\nSubject: {msg['Subject']}"
    paras = ['<div class="emlhdr">' + H.escape(hdr).replace('\n', '<br>') + '</div>']
    for block in re.split(r'\n{2,}', body_txt):
        b = block.strip()
        if not b:
            continue
        e = re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', H.escape(b).replace('\n', '<br>'))
        cls = 'emlhdr' if re.match(r'^(From|Sent|To|Subject):', b) else None
        paras.append(f'<div class="emlhdr">{e}</div>' if cls else f'<p>{e}</p>')
    return '\n'.join(paras), hdr + '\n' + body_txt


def xlsx_to_html(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    parts, plain = [], []
    for ws in wb.worksheets:
        parts.append(f'<h3>Sheet: {H.escape(ws.title)}</h3>')
        rows, first = [], True
        for row in ws.iter_rows(values_only=True):
            if all(v is None or str(v).strip() == '' for v in row):
                continue
            cells = ['' if v is None else
                     (f'{v:,.2f}'.rstrip('0').rstrip('.') if isinstance(v, float) else str(v))
                     for v in row]
            plain.append(' | '.join(cells))
            tc = 'th' if first else 'td'
            rows.append('<tr>' + ''.join(f'<{tc}>{H.escape(c)}</{tc}>' for c in cells) + '</tr>')
            first = False
        parts.append('<div class="tblwrap"><table>' + ''.join(rows) + '</table></div>')
    return '\n'.join(parts), '\n'.join(plain)


CONVERTERS = {'.docx': docx_to_html, '.eml': eml_to_html, '.xlsx': xlsx_to_html}

STOP_NUMS = {'2010', '2020', '2021', '2022', '2023', '2024', '2025', '2026', '100', '1000'}


def tokens_for(rubric):
    """Distinctive tokens from a rubric: figures, plus phrases the rubric itself quotes."""
    txt = (rubric.get('title') or '') + ' ' + (rubric.get('match_criteria') or '')
    toks = set()
    for m in re.finditer(r'\$?(\d{1,2}[,.]\d{2,3}|\d{3,4}(?:,\d{3})?|\d+\.\d+)', txt):
        num = m.group(1)
        if num.replace(',', '') in STOP_NUMS or num.startswith('0'):
            continue
        if '.' in num or ',' in num or len(num.replace(',', '')) >= 3:
            toks.add(num)
    # quoted phrases ('moderately concentrated', "hell-or-high-water") are the
    # rubric author flagging exact expected wording - strong link signals
    for m in re.finditer(r"['\"]([A-Za-z][A-Za-z0-9 \-]{5,60})['\"]", txt):
        toks.add(m.group(1).lower())
    return toks


def link_rubrics(rubrics, docs):
    plain = {d['id']: d['plain'].lower() for d in docs}
    links = {}
    for r in rubrics:
        hits = []
        for d in docs:
            matched = set()
            for t in tokens_for(r):
                variants = {t, t.replace(',', '')}
                if re.fullmatch(r'\d{4}', t):
                    variants.add(f'{t[0]},{t[1:]}')
                if any(v.lower() in plain[d['id']] for v in variants):
                    matched.add(t)
            if matched:
                hits.append({'doc': d['id'], 'tokens': sorted(matched)[:4], 'n': len(matched)})
        hits.sort(key=lambda h: -h['n'])
        links[r['id']] = hits[:5]
    return links


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='out')
    args = ap.parse_args()
    out = pathlib.Path(args.out)
    ex = out / 'extracted'
    docs_dir = out / 'source-docs'
    docs_dir.mkdir(exist_ok=True)

    urls = json.loads((ex / 'documents.json').read_text())
    page = json.loads((out / 'page-data.json').read_text())
    deliverables = page.get('outputs')
    if not isinstance(deliverables, dict):
        deliverables = {}
    # accept plain URL strings or {url: ...} dicts; drop anything else
    clean_urls = []
    for u in urls if isinstance(urls, list) else []:
        if isinstance(u, dict):
            u = u.get('url')
        if isinstance(u, str) and u:
            clean_urls.append(u)
        else:
            print(f'WARN: skipping non-URL document entry: {u!r}')

    def dest_name(url):
        base = url.rsplit('/', 1)[-1].split('?')[0] or 'document'
        base = base.replace('/', '_').replace('..', '_')
        if [c.rsplit('/', 1)[-1].split('?')[0] for c in clean_urls].count(base) > 1:
            base = hashlib.sha1(url.encode()).hexdigest()[:8] + '-' + base
        return base

    to_fetch = [(u, docs_dir / dest_name(u), 'source') for u in clean_urls]
    for fname, url in deliverables.items():
        if isinstance(url, str) and isinstance(fname, str):
            to_fetch.append((url, docs_dir / ('deliverable-' + fname.replace('/', '_')), 'deliverable'))

    docs = []
    for url, dest, kind in to_fetch:
        try:
            if not dest.exists():
                urllib.request.urlretrieve(url, dest)
            conv = CONVERTERS.get(dest.suffix.lower())
            if not conv:
                print(f'skip (no converter): {dest.name}')
                continue
            body_html, plain = conv(dest)
            doc_id = re.sub(r'\.\w+$', '', dest.name)
            title = ('FINAL DELIVERABLE - ' + dest.name.replace('deliverable-', '')
                     if kind == 'deliverable' else dest.name)
            docs.append({'id': doc_id, 'file': dest.name,
                         'kind': 'deliverable' if kind == 'deliverable' else dest.suffix.lstrip('.'),
                         'title': title, 'html': body_html, 'plain': plain})
            print(f'converted {dest.name} ({len(body_html)} chars)')
        except Exception as e:
            print(f'WARN: {dest.name}: {e}')

    json.dump(docs, open(ex / 'source-docs.json', 'w'))
    links = link_rubrics(page.get('rubrics', []), docs)
    json.dump(links, open(ex / 'rubric-doc-links.json', 'w'))
    linked = sum(1 for v in links.values() if v)
    print(f'{len(docs)} documents converted; {linked}/{len(links)} rubrics linked to sources')


if __name__ == '__main__':
    main()
